"""
Build the CRO synthetic Word samples (all content invented).

    python3 build_word_samples.py      # writes into ./out, appends to manifest
"""
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, Emu, Mm

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
MARK = "SAMPLE – synthetic test document – not a real agreement"
manifest_path = os.path.join(OUT, "manifest_pdf.json")
manifest = json.load(open(manifest_path, encoding="utf-8"))

RIGHT_OF_LABEL = {"percentX": 0.42, "percentY": -0.012, "percentWidth": 0.22, "percentHeight": 0.045}
DATE_OFFSET = {"percentX": 0.42, "percentY": -0.008, "percentWidth": 0.18, "percentHeight": 0.022}


def base_doc(title):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    f = sec.footer.paragraphs[0]
    f.text = MARK
    f.runs[0].font.size = Pt(7)
    d.add_heading(title, level=1)
    return d


def rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(parse_xml('<w:bidi xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
    for r in paragraph.runs:
        rPr = r._r.get_or_add_rPr()
        rPr.append(parse_xml('<w:rtl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>'))
        r.font.name = "Arial"
        rPr.rFonts.set(qn("w:cs"), "Arial")


# ---------------------------------------------------------------- 10 Bilingual tenancy addendum
d = base_doc("Tenancy Contract Addendum / ملحق عقد الإيجار")
d.add_paragraph("Landlord: Palm Crescent Properties LLC (fictional). Tenant: Yusuf Al Mansoori (fictional). "
                "Property: Unit 7, Sample Residence, Dubai (fictional).")
p = d.add_paragraph("المؤجر: شركة نخلة الهلال للعقارات (اسم وهمي). المستأجر: يوسف المنصوري (اسم وهمي).")
rtl(p)
d.add_heading("1. Rent revision / تعديل الإيجار", level=2)
d.add_paragraph("With effect from 1 January 2027 (sample) the annual rent is revised to AED 96,000 payable in four cheques.")
p = d.add_paragraph("اعتبارًا من ١ يناير ٢٠٢٧ يتم تعديل الإيجار السنوي إلى ٩٦٬٠٠٠ درهم تدفع على أربع دفعات.")
rtl(p)
d.add_heading("2. Other terms / شروط أخرى", level=2)
d.add_paragraph("All other terms of the tenancy contract remain unchanged.")
p = d.add_paragraph("تبقى جميع شروط عقد الإيجار الأخرى دون تغيير.")
rtl(p)
d.add_paragraph()
t = d.add_table(rows=4, cols=2)
t.style = "Table Grid"
cells = [("Landlord Signature / توقيع المؤجر:", ""), ("Landlord Date / التاريخ:", ""),
         ("Tenant Signature / توقيع المستأجر:", ""), ("Tenant Date / التاريخ:", "")]
for i, (a, b) in enumerate(cells):
    t.cell(i, 0).text = a
    t.rows[i].height = Mm(16 if "Signature" in a else 9)
d.save(os.path.join(OUT, "10_bilingual_tenancy_addendum.docx"))
manifest.append({"file": "10_bilingual_tenancy_addendum.docx", "title": "Bilingual tenancy addendum (Word)",
                 "purpose": "Word → PDF conversion by Circularo, Arabic RTL text, English part of bilingual label as anchor",
                 "placement": "anchors",
                 "recipients": [{"role": "landlord", "kind": "internal", "purpose": "sign", "order": 1},
                                {"role": "tenant", "kind": "external", "purpose": "sign", "order": 2, "protection": "email_otp"}],
                 "anchors": [{"id": "SIG_LANDLORD", "text": "Landlord Signature", "role": "landlord", "type": "signature", "offset": RIGHT_OF_LABEL},
                             {"id": "DATE_LANDLORD", "text": "Landlord Date", "role": "landlord", "type": "annotation", "subtype": "annotation_date", "offset": DATE_OFFSET},
                             {"id": "SIG_TENANT", "text": "Tenant Signature", "role": "tenant", "type": "signature", "offset": RIGHT_OF_LABEL},
                             {"id": "DATE_TENANT", "text": "Tenant Date", "role": "tenant", "type": "annotation", "subtype": "annotation_date", "offset": DATE_OFFSET}],
                 "notes": "Checks Circularo's Office conversion (or TTA-side conversion) and that anchors survive it. Also try an Arabic anchor (توقيع المستأجر)."})

# ---------------------------------------------------------------- 11 Dynamic MS Word eSigning template
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_shape_id = [100]


def rect_with_text(paragraph, text, w_mm=60, h_mm=16):
    """Insert an inline DrawingML rectangle containing `text` (small grey font)."""
    _shape_id[0] += 1
    cx, cy = int(Emu(Mm(w_mm))), int(Emu(Mm(h_mm)))
    sid = _shape_id[0]
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    xml = f'''
<w:r xmlns:w="{W}">
  <w:drawing>
    <wp:inline xmlns:wp="{WP}" distT="0" distB="0" distL="0" distR="0">
      <wp:extent cx="{cx}" cy="{cy}"/>
      <wp:docPr id="{sid}" name="Signature field {sid}"/>
      <a:graphic xmlns:a="{A}">
        <a:graphicData uri="{WPS}">
          <wps:wsp xmlns:wps="{WPS}">
            <wps:cNvSpPr/>
            <wps:spPr>
              <a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
              <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
              <a:noFill/>
              <a:ln w="6350"><a:solidFill><a:srgbClr val="BFBFBF"/></a:solidFill></a:ln>
            </wps:spPr>
            <wps:txbx>
              <w:txbxContent>
                <w:p><w:r><w:rPr><w:color w:val="A6A6A6"/><w:sz w:val="2"/><w:szCs w:val="2"/></w:rPr><w:t xml:space="preserve">{text}</w:t></w:r></w:p>
              </w:txbxContent>
            </wps:txbx>
            <wps:bodyPr rot="0" lIns="18000" tIns="18000" rIns="18000" bIns="18000" anchor="t"/>
          </wps:wsp>
        </a:graphicData>
      </a:graphic>
    </wp:inline>
  </w:drawing>
</w:r>'''
    paragraph._p.append(parse_xml(xml))


def field(obj):
    # Circularo requires no whitespace inside the JSON.
    return "signature_field_start:" + json.dumps(obj, separators=(",", ":"), ensure_ascii=False) + ":signature_field_end"


d = base_doc("Service Order Form – dynamic template (sample)")
d.add_paragraph("Customer: Cedar Point Hospitality LLC (fictional). Service: quarterly HVAC maintenance (sample).")
for i in range(4):
    d.add_paragraph("The service provider shall perform the scheduled maintenance visits listed in the annex, "
                    "record each visit and report defects within two working days. (Clause %d, sample text.)" % (i + 1))
d.add_heading("Customer", level=2)
p = d.add_paragraph("Full name: ")
rect_with_text(p, field({"type": "annotation", "subtype": "annotation_full_name", "userRoleId": "customer", "userRoleName": "Customer", "required": True}), 70, 7)
p = d.add_paragraph("Signature: ")
rect_with_text(p, field({"type": "signature", "userRoleId": "customer", "userRoleName": "Customer", "required": True, "timestamp": True}))
p = d.add_paragraph("Date: ")
rect_with_text(p, field({"type": "annotation", "subtype": "annotation_date", "userRoleId": "customer", "userRoleName": "Customer"}), 40, 7)
p = d.add_paragraph("Service frequency: ")
rect_with_text(p, field({"type": "annotation", "subtype": "annotation_dropdown", "userRoleId": "customer", "userRoleName": "Customer",
                         "values": ["Monthly", "Quarterly", "Annually"]}), 45, 7)
p = d.add_paragraph("I accept the annex: ")
rect_with_text(p, field({"type": "annotation", "subtype": "annotation_check", "userRoleId": "customer", "userRoleName": "Customer", "required": True}), 10, 10)
d.add_heading("Service provider", level=2)
p = d.add_paragraph("Signature: ")
rect_with_text(p, field({"type": "signature", "userRoleId": "provider", "userRoleName": "Provider", "required": True}))
p = d.add_paragraph("Company stamp: ")
rect_with_text(p, field({"type": "stamp", "userRoleId": "provider", "userRoleName": "Provider"}), 35, 25)
d.save(os.path.join(OUT, "11_dynamic_word_template.docx"))
manifest.append({"file": "11_dynamic_word_template.docx", "title": "Dynamic MS Word eSigning template",
                 "purpose": "Circularo's own field notation inside Word rectangle shapes (roles, dropdown, checkbox, stamp, timestamp)",
                 "placement": "dynamic-word-template",
                 "recipients": [{"role": "customer", "kind": "external", "purpose": "sign", "order": 1, "protection": "email_otp"},
                                {"role": "provider", "kind": "internal", "purpose": "sign", "order": 2}],
                 "anchors": [],
                 "notes": "Upload as a template in Circularo (admin) or upload via API and inspect signFields/formFields. If Circularo does not "
                          "detect these shapes, re-save the file once in Microsoft Word (shapes were generated programmatically)."})

# ---------------------------------------------------------------- 12 Ad hoc Word letter
d = base_doc("Letter of Confirmation (sample)")
d.add_paragraph("To whom it may concern,")
for i in range(5):
    d.add_paragraph("We confirm that the sample services described in our earlier correspondence were delivered as agreed "
                    "and that no outstanding amounts are due. (Paragraph %d, sample text.)" % (i + 1))
d.add_paragraph("Yours faithfully,")
d.add_paragraph("Falcon Ridge Analytics LLC (fictional)")
d.save(os.path.join(OUT, "12_adhoc_letter.docx"))
manifest.append({"file": "12_adhoc_letter.docx", "title": "Ad hoc Word letter", "purpose": "Arbitrary Word upload with no anchors",
                 "placement": "adhoc",
                 "recipients": [{"role": "signatory", "kind": "internal", "purpose": "sign", "order": 1}],
                 "anchors": [], "notes": "Tests conversion plus ad hoc placement together."})

with open(os.path.join(OUT, "manifest.json"), "w", encoding="utf-8") as f:
    json.dump({"generated": "2026-10-15", "note": "Synthetic test documents for the CRO Circularo integration",
               "documents": manifest}, f, indent=2, ensure_ascii=False)
os.remove(manifest_path)
print("Total samples:", len(manifest))

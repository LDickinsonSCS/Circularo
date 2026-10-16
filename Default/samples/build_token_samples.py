"""
Ad hoc samples marked with {{sign:n}} tokens (n = recipient number in CRO_Send / CRO_REQUEST).
Synthetic content only.   python3 build_token_samples.py   # writes into ./out
"""
import os
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
import docx
from docx.shared import Pt, RGBColor

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
os.makedirs(OUT, exist_ok=True)
MARK = "SAMPLE – synthetic test document – not a real agreement"
TEXT = [
    "This letter confirms the arrangements discussed for the delivery of the pilot project described below.",
    "The supplier will provide the services from 1 November for a period of three months, after which the parties will review the results.",
    "Fees, reporting and acceptance follow the attached schedule. Either party may end the pilot with fourteen days' notice in writing.",
    "Please sign where marked to confirm agreement. Token placement test: each {{sign:n}} marker becomes the signature field of recipient n.",
]

# PDF: two pages, tokens for recipients 1 and 2 on page 2, a third (date-style) token omitted on purpose
p = os.path.join(OUT, "13_adhoc_tokens.pdf")
c = canvas.Canvas(p, pagesize=A4)
W, H = A4
for page in (1, 2):
    c.setFont("Helvetica-Bold", 15)
    c.drawString(20 * mm, H - 25 * mm, "Pilot project confirmation (token sample)" if page == 1 else "Signatures")
    c.setFont("Helvetica", 10)
    y = H - 40 * mm
    if page == 1:
        for t in TEXT * 3:
            c.drawString(20 * mm, y, t[:105]); y -= 6 * mm
            if len(t) > 105: c.drawString(20 * mm, y, t[105:]); y -= 6 * mm
            y -= 3 * mm
    else:
        for i, (who, org) in enumerate([("Amira Haddad", "Northwind Trading LLC (sample)"), ("Tom Becker", "Contoso Services FZE (sample)")]):
            top = H - 60 * mm - i * 55 * mm
            c.setFont("Helvetica-Bold", 10); c.drawString(20 * mm, top, f"Signed for {org}")
            c.setFont("Helvetica", 8); c.setFillColor(colors.HexColor("#9AA4B2"))
            c.drawString(20 * mm, top - 14 * mm, "{{sign:%d}}" % (i + 1))
            c.setFillColor(colors.black)
            c.line(20 * mm, top - 16 * mm, 95 * mm, top - 16 * mm)
            c.setFont("Helvetica", 10); c.drawString(20 * mm, top - 22 * mm, f"Name: {who}")
            c.drawString(20 * mm, top - 28 * mm, "Date: ____________________")
    c.setFont("Helvetica", 7); c.setFillColor(colors.grey)
    c.drawString(20 * mm, 10 * mm, MARK); c.drawRightString(190 * mm, 10 * mm, f"Page {page}")
    c.setFillColor(colors.black)
    c.showPage()
c.save()

# Word: tokens in white text (invisible once signed), recipient 1 and 2
d = docx.Document()
d.add_heading("Service order acceptance (token sample, Word)", 1)
for t in TEXT:
    d.add_paragraph(t)
for i, (who, org) in enumerate([("Amira Haddad", "Northwind Trading LLC (sample)"), ("Tom Becker", "Contoso Services FZE (sample)")]):
    d.add_paragraph().add_run(f"Signed for {org}").bold = True
    r = d.add_paragraph().add_run("{{sign:%d}}" % (i + 1)); r.font.size = Pt(8); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    d.add_paragraph("______________________________")
    d.add_paragraph(f"Name: {who}")
    d.add_paragraph("")
sec = d.sections[0]; sec.footer.paragraphs[0].text = MARK
d.save(os.path.join(OUT, "14_adhoc_tokens.docx"))
print("written 13_adhoc_tokens.pdf, 14_adhoc_tokens.docx")
